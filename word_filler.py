"""
Compilazione moduli Word (.docx in uscita).

Gemello di `pdf_filler.py` per i moduli che esistono anche in formato Word.
Riceve i dati di una pratica + cliente e restituisce i byte di un .docx compilato.
Non conosce la UI: in ingresso dict, in uscita bytes.

Due motori, scelti dal registro WORD_TEMPLATES:

  "celle"      → il template è un modulo ufficiale già impaginato (celle fisse,
                 nessun segnaposto): si scrive per coordinate (riga, colonna).
                 Oggi lo usa la sola "Prescrizione HBG".

  "segnaposto" → il template è un documento nostro (carta intestata) con
                 segnaposto {{chiave}} nel testo, nelle tabelle e nelle
                 intestazioni. Le righe dei codici si ottengono duplicando la
                 riga di tabella che contiene {{codice}}.
                 Lo usa la "Relazione tecnica".
"""
import copy
import io
import os

from docx import Document

from pdf_filler import dati_canonici, _fmt_qta

TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             "assets", "word-templates")

# Registro dei moduli disponibili in Word.
#   gemello_pdf: True → esiste anche come PDF con lo stesso id, quindi nella
#   scheda pratica compare come pulsante "Word" accanto a "Genera".
WORD_TEMPLATES = {
    "prescrizione-hbg": {
        "label": "Prescrizione HBG",
        "file": "Prescrizione HBG.docx",
        "motore": "celle",
        "gemello_pdf": True,
    },
    "relazione-tecnica": {
        "label": "Relazione tecnica",
        "file": "Relazione Tecnica.docx",
        "motore": "segnaposto",
        "richiede_cliente": True,
    },
}


def moduli_word_gemelli() -> set:
    """Id dei moduli Word che affiancano un PDF omonimo (pulsante 'Word')."""
    return {tid for tid, m in WORD_TEMPLATES.items() if m.get("gemello_pdf")}


def template_disponibile(template_id: str) -> bool:
    """True se il file .docx del modulo è presente in assets/word-templates."""
    tpl = WORD_TEMPLATES.get(template_id)
    return bool(tpl) and os.path.isfile(os.path.join(TEMPLATES_DIR, tpl["file"]))


# ── Utilità comuni ────────────────────────────────────────────────────────────

def _scrivi_run(run, testo: str) -> None:
    """Sostituisce il testo di un run mantenendone la formattazione. Gli a-capo
    del testo diventano interruzioni di riga vere (il campo posturale è un
    textarea, quindi può contenere più righe)."""
    parti = (testo or "").split("\n")
    run.text = parti[0]
    for p in parti[1:]:
        run.add_break()
        run.add_text(p)


def _svuota_paragrafi_extra(cella) -> None:
    for par in cella.paragraphs[1:]:
        for run in par.runs:
            run.text = ""


# ── Motore 1: scrittura per coordinate (moduli ufficiali) ─────────────────────

# Mappa del template HBG: la tabella dati è la seconda del documento (indice 1).
# Le celle sono individuate da (riga, colonna) perché il modulo usa celle unite
# senza segnaposto testuali. Indici rilevati dal template in assets/word-templates.
_HBG_TABELLA = 1
_HBG_ANAGRAFICA = {          # chiave dati comuni → (riga, colonna)
    "nome":       (0, 3),    # Cognome e nome
    "nato_data":  (0, 18),   # nato il
    "nato_luogo": (2, 1),    # a
    "comune":     (2, 12),   # residente a
    "provincia":  (2, 21),   # Prov.
    "cap":        (4, 2),    # C.A.P.
    "via":        (4, 6),    # Via
    "telefono":   (4, 17),   # Tel.
}
_HBG_DIAGNOSI = (8, 10)
_HBG_SIGNIFICATO = (23, 8)
_HBG_RIGHE_DA = 12           # prima riga ausili
_HBG_RIGHE_A = 22            # ultima riga ausili (inclusa)
_HBG_COL_DESCRIZIONE = 0
_HBG_COL_CODICE = 20


def _scrivi_cella(tabella, riga: int, colonna: int, testo: str) -> None:
    """Scrive nella cella mantenendo la formattazione del template: riusa il primo
    run del primo paragrafo (che porta font e corpo del modulo) e svuota il resto."""
    if riga >= len(tabella.rows):
        return
    cella = tabella.rows[riga].cells[colonna]
    testo = testo or ""
    for i, par in enumerate(cella.paragraphs):
        if i == 0:
            if par.runs:
                _scrivi_run(par.runs[0], testo)
                for run in par.runs[1:]:
                    run.text = ""
            elif testo:
                _scrivi_run(par.add_run(""), testo)
        else:
            for run in par.runs:
                run.text = ""


def _descrizione_con_qta(riga: dict) -> str:
    """Dicitura dell'ausilio. Il modulo Word non ha una colonna quantità: quando
    è maggiore di 1 la si accoda alla descrizione, altrimenti si omette."""
    desc = (riga.get("descrizione") or "").strip()
    try:
        qta = float(riga.get("qta") or 0)
    except (TypeError, ValueError):
        qta = 0
    if desc and qta > 1:
        return f"{desc} (q.tà {_fmt_qta(qta)})"
    return desc


def _compila_hbg(doc, pratica: dict, cliente: dict, righe: list) -> None:
    tabella = doc.tables[_HBG_TABELLA]
    dati = dati_canonici(pratica, cliente)

    for chiave, (riga, colonna) in _HBG_ANAGRAFICA.items():
        _scrivi_cella(tabella, riga, colonna, dati.get(chiave, ""))
    _scrivi_cella(tabella, *_HBG_DIAGNOSI, dati.get("diagnosi", ""))
    _scrivi_cella(tabella, *_HBG_SIGNIFICATO, dati.get("sign_terapeutico", ""))

    # Righe ausili: descrizione (+ q.tà) e codice ISO, una riga di tabella ciascuna.
    max_righe = _HBG_RIGHE_A - _HBG_RIGHE_DA + 1
    for i, riga in enumerate((righe or [])[:max_righe]):
        r = _HBG_RIGHE_DA + i
        _scrivi_cella(tabella, r, _HBG_COL_DESCRIZIONE, _descrizione_con_qta(riga))
        _scrivi_cella(tabella, r, _HBG_COL_CODICE, (riga.get("codice_iso") or "").strip())


# ── Motore 2: sostituzione segnaposto {{chiave}} ──────────────────────────────

# Segnaposto della riga codici: la riga di tabella che li contiene viene
# duplicata una volta per ogni codice della pratica.
_SEGNAPOSTO_RIGA = ("codice", "descrizione", "qta")


def valori_segnaposto(pratica: dict, cliente: dict) -> dict:
    """Vocabolario dei segnaposto documentato all'utente. Poggia su
    `dati_canonici`, così i documenti Word e i moduli PDF dicono le stesse cose."""
    d = dati_canonici(pratica, cliente)
    pratica = pratica or {}
    return {
        # Paziente
        "paziente":          d["nome"],
        "nato_il":           d["nato_data"],
        "nato_a":            d["nato_luogo"],
        "residenza":         d["via"],
        "citta":             d["citta"],
        "comune":            d["comune"],
        "provincia":         d["provincia"],
        "cap":               d["cap"],
        "codice_fiscale":    d["cf"],
        "telefono":          d["telefono"],
        "email":             d["email"],
        # Pratica
        "asl":               d["asl"],
        "centro":            d["centro"],
        "medico":            d["medico_struttura"],
        "data_prescrizione": d["data_prescrizione"],
        "numero_pratica":    (pratica.get("numero_pratica") or "").strip(),
        "ausilio":           d["ausilio"],
        # Contenuti
        "diagnosi":                d["diagnosi"],
        "descrizione_posturale":   (pratica.get("descrizione_posturale") or "").strip(),
        "significato_terapeutico": d["sign_terapeutico"],
        # Data del documento
        "oggi":       d["oggi"],
        "luogo_data": d["luogo_data"],
    }


def _applica(testo: str, valori: dict) -> str:
    """Sostituisce i segnaposto noti. Quelli sconosciuti restano visibili di
    proposito: così un refuso nel template si vede subito nel documento."""
    for chiave, valore in valori.items():
        segnaposto = "{{" + chiave + "}}"
        if segnaposto in testo:
            testo = testo.replace(segnaposto, valore or "")
    return testo


def _sostituisci_paragrafo(par, valori: dict) -> None:
    """Sostituisce i segnaposto di un paragrafo.

    Prima prova run per run: se il segnaposto sta tutto dentro un run, la
    formattazione mista del paragrafo ('Diagnosi:' in grassetto + valore normale)
    resta intatta. Solo se Word ha spezzato il segnaposto su più run si ricade
    sull'unificazione nel primo run."""
    if not par.runs:
        return
    for run in par.runs:
        if "{{" in run.text:
            nuovo = _applica(run.text, valori)
            if nuovo != run.text:
                _scrivi_run(run, nuovo)
    testo = "".join(r.text for r in par.runs)
    if "{{" in testo:
        nuovo = _applica(testo, valori)
        if nuovo != testo:
            _scrivi_run(par.runs[0], nuovo)
            for run in par.runs[1:]:
                run.text = ""


def _sostituisci_in_tabella(tabella, valori: dict) -> None:
    for riga in tabella.rows:
        for cella in riga.cells:
            for par in cella.paragraphs:
                _sostituisci_paragrafo(par, valori)
            for annidata in cella.tables:      # tabelle dentro le celle
                _sostituisci_in_tabella(annidata, valori)


def _sostituisci_in_contenitore(contenitore, valori: dict) -> None:
    for par in contenitore.paragraphs:
        _sostituisci_paragrafo(par, valori)
    for tabella in contenitore.tables:
        _sostituisci_in_tabella(tabella, valori)


def _sostituisci_ovunque(doc, valori: dict) -> None:
    """Corpo del documento + intestazioni e piè di pagina (dove sta la carta
    intestata), comprese quelle di prima pagina e pagine pari."""
    _sostituisci_in_contenitore(doc, valori)
    for sezione in doc.sections:
        for nome in ("header", "footer", "first_page_header", "first_page_footer",
                     "even_page_header", "even_page_footer"):
            parte = getattr(sezione, nome, None)
            if parte is not None:
                _sostituisci_in_contenitore(parte, valori)


def _valori_riga(riga: dict) -> dict:
    try:
        qta = float(riga.get("qta") or 0)
    except (TypeError, ValueError):
        qta = 0
    return {
        "codice":      (riga.get("codice_iso") or "").strip(),
        "descrizione": (riga.get("descrizione") or "").strip(),
        "qta":         _fmt_qta(qta) if qta else "",
    }


def _espandi_righe_codici(doc, righe: list) -> bool:
    """Trova la riga-modello dei codici (quella che contiene {{codice}}) e la
    duplica una volta per ogni codice della pratica, poi la elimina. Se la
    pratica non ha codici resta la sola intestazione della tabella."""
    attesi = tuple("{{" + s + "}}" for s in _SEGNAPOSTO_RIGA)
    for tabella in doc.tables:
        indice = None
        for i, riga in enumerate(tabella.rows):
            testo = "".join(c.text for c in riga.cells)
            if any(s in testo for s in attesi):
                indice = i
                break
        if indice is None:
            continue

        modello_tr = tabella.rows[indice]._tr
        for _ in (righe or []):
            modello_tr.addprevious(copy.deepcopy(modello_tr))
        # Le copie ora occupano le posizioni indice … indice+n-1; `tabella.rows`
        # rilegge l'XML a ogni accesso, quindi bastano gli indici.
        for i, riga in enumerate(righe or []):
            for cella in tabella.rows[indice + i].cells:
                for par in cella.paragraphs:
                    _sostituisci_paragrafo(par, _valori_riga(riga))
        modello_tr.getparent().remove(modello_tr)
        return True
    return False


def _compila_segnaposto(doc, pratica: dict, cliente: dict, righe: list) -> None:
    # Prima le righe codici (duplica la riga-modello), poi il resto del documento:
    # così i segnaposto di riga non vengono azzerati prima di essere clonati.
    _espandi_righe_codici(doc, righe or [])
    _sostituisci_ovunque(doc, valori_segnaposto(pratica, cliente))


# ── Ingresso pubblico ─────────────────────────────────────────────────────────

_MOTORI = {
    "celle":      _compila_hbg,
    "segnaposto": _compila_segnaposto,
}


def compila_docx(template_id: str, pratica: dict, cliente: dict, righe: list = None) -> bytes:
    """Compila il modulo Word e restituisce i byte del .docx."""
    tpl = WORD_TEMPLATES.get(template_id)
    if not tpl:
        raise ValueError(f"Modulo Word non disponibile: {template_id}")

    path = os.path.join(TEMPLATES_DIR, tpl["file"])
    if not os.path.isfile(path):
        raise FileNotFoundError(f"File template mancante: {path}")

    doc = Document(path)
    _MOTORI[tpl.get("motore", "segnaposto")](doc, pratica, cliente, righe or [])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
